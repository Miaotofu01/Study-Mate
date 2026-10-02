"""附件文档解析：字节进、文本出（对照 DeepTutor document_extractor 的确定性管线）。

文本类多编码回退解码；pdf→pymupdf（缺失回退 pypdf，皆缺报错占位）；
docx→python-docx；xlsx→openpyxl；pptx→python-pptx；epub→标准库 ZIP+OPF。
解析失败抛异常，由调用方转为「[附件 <文件名> 解析失败：<原因>]」占位。
"""
from __future__ import annotations

import io
import posixpath
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from xml.etree import ElementTree

MAX_CHARS_PER_DOC = 20000

TEXT_LIKE_EXTENSIONS = frozenset(
    {
        ".md", ".markdown", ".txt", ".csv", ".json", ".xml", ".yaml", ".yml",
        ".toml", ".ini", ".svg",
        ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".scss", ".html", ".htm",
        ".c", ".h", ".cpp", ".hpp", ".java", ".go", ".rs", ".rb", ".php",
        ".sh", ".bat", ".sql", ".ipynb",
    }
)

_DECODE_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")
_EPUB_CONTENT_EXTENSIONS = (".xhtml", ".html", ".htm")
_EPUB_CONTAINER = "META-INF/container.xml"


def _decode_text(data: bytes) -> str:
    for encoding in _DECODE_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class _HTMLTextParser(HTMLParser):
    _BLOCK_TAGS = frozenset(
        {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "section"}
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "br":
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif not self._skip and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html_text(raw: bytes) -> str:
    parser = _HTMLTextParser()
    parser.feed(_decode_text(raw))
    parser.close()
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in "".join(parser.parts).splitlines()]
    return "\n".join(line for line in lines if line)


def _extract_pdf(data: bytes, filename: str) -> str:
    first_error: Exception | None = None
    try:
        import fitz
    except ImportError:
        fitz = None
    if fitz is not None:
        try:
            with fitz.open(stream=data, filetype="pdf") as doc:
                if doc.is_encrypted and not doc.authenticate(""):
                    raise ValueError("PDF 已加密，无法解析")
                pages = [
                    f"--- Page {number} ---\n{page.get_text() or ''}"
                    for number, page in enumerate(doc, 1)
                ]
            return "\n\n".join(pages)
        except Exception as exc:  # noqa: BLE001 - 转入 pypdf 回退
            first_error = exc
    try:
        from pypdf import PdfReader
    except ImportError:
        if first_error is not None:
            raise RuntimeError(f"PDF 解析失败（{first_error}）") from None
        raise RuntimeError("缺少 PDF 解析库，请安装 pymupdf 或 pypdf") from None
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [
            f"--- Page {number} ---\n{page.extract_text() or ''}"
            for number, page in enumerate(reader.pages, 1)
        ]
        return "\n\n".join(pages)
    except Exception as exc:  # noqa: BLE001 - 统一转运行时错误
        raise RuntimeError(f"PDF 解析失败（{exc}）") from None


def _extract_docx(data: bytes, filename: str) -> str:
    try:
        import docx
    except ImportError:
        raise RuntimeError("缺少 python-docx，无法解析 docx") from None
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"docx 解析失败（{exc}）") from None
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes, filename: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("缺少 openpyxl，无法解析 xlsx") from None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"xlsx 解析失败（{exc}）") from None
    sheets: list[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                row_text = "\t".join("" if cell is None else str(cell) for cell in row)
                if row_text.strip():
                    rows.append(row_text)
            if rows:
                sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def _collect_shape_text(shapes: Any, out: list[str]) -> None:
    for shape in shapes:
        sub_shapes = getattr(shape, "shapes", None)
        if sub_shapes is not None:
            _collect_shape_text(sub_shapes, out)
            continue
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = "\t".join(cell for cell in cells if cell)
                if line:
                    out.append(line)
            continue
        text = getattr(shape, "text", "")
        if text:
            out.append(text)


def _extract_pptx(data: bytes, filename: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("缺少 python-pptx，无法解析 pptx") from None
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"pptx 解析失败（{exc}）") from None
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, 1):
        texts: list[str] = []
        _collect_shape_text(slide.shapes, texts)
        if texts:
            slides.append(f"--- Slide {number} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _local_name(tag: Any) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) and "}" in tag else (tag or "")


def _extract_epub(data: bytes, filename: str) -> str:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"epub 解析失败（{exc}）") from None
    with archive:
        names = set(archive.namelist())
        opf_path = ""
        if _EPUB_CONTAINER in names:
            try:
                root = ElementTree.fromstring(archive.read(_EPUB_CONTAINER))
                for node in root.iter():
                    if _local_name(node.tag) == "rootfile":
                        opf_path = node.get("full-path") or ""
                        break
            except ElementTree.ParseError:
                opf_path = ""
        members: list[str] = []
        if opf_path and opf_path in names:
            try:
                opf_root = ElementTree.fromstring(archive.read(opf_path))
                manifest: dict[str, str] = {}
                spine: list[str] = []
                opf_dir = posixpath.dirname(opf_path)
                for node in opf_root.iter():
                    name = _local_name(node.tag)
                    if name == "item" and node.get("id") and node.get("href"):
                        manifest[node.get("id")] = node.get("href") or ""
                    elif name == "itemref" and node.get("idref"):
                        spine.append(node.get("idref") or "")
                for idref in spine:
                    href = manifest.get(idref)
                    if not href:
                        continue
                    member = posixpath.normpath(posixpath.join(opf_dir, href.split("#", 1)[0]))
                    if (
                        member in names
                        and member not in members
                        and member.lower().endswith(_EPUB_CONTENT_EXTENSIONS)
                    ):
                        members.append(member)
            except ElementTree.ParseError:
                members = []
        if not members:
            members = sorted(name for name in names if name.lower().endswith(_EPUB_CONTENT_EXTENSIONS))
        texts: list[str] = []
        for member in members:
            try:
                texts.append(_html_text(archive.read(member)))
            except Exception:  # noqa: BLE001 - 单章失败跳过
                continue
        return "\n\n".join(text for text in texts if text)


def extract_text(filename: str, data: bytes, max_chars: int = MAX_CHARS_PER_DOC) -> str:
    """提取单个附件文本；失败抛异常（调用方转占位），输出超限截断。"""
    if not data:
        raise ValueError("文件内容为空")
    ext = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(data, filename)
    elif ext == ".docx":
        text = _extract_docx(data, filename)
    elif ext == ".xlsx":
        text = _extract_xlsx(data, filename)
    elif ext == ".pptx":
        text = _extract_pptx(data, filename)
    elif ext == ".epub":
        text = _extract_epub(data, filename)
    elif ext in TEXT_LIKE_EXTENSIONS:
        text = _decode_text(data)
    else:
        raise ValueError(f"不支持的文件类型：{ext or '未知'}")
    if not text.strip():
        raise ValueError("未解析出文本内容")
    if len(text) > max_chars:
        text = f"{text[:max_chars]}\n…（已截断，原文共 {len(text)} 字符）"
    return text
